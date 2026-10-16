import pandas as pd
from datetime import datetime
from io import BytesIO
from pathlib import Path
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent.parent
TEMPLATE_PATH = BASE_DIR / "templates" / "TIME_&_LABOR_FORM_WE.xlsx"


def CAK(uploaded_file):
    
    # 1. Read raw sheet
    df_raw = pd.read_excel(uploaded_file, header=None)

    # 2. Extract structured records (Employee ID, Name, Action, Date, Hours)
    records = []
    emp_id, emp_name, curr_date = None, None, None

    for _, row in df_raw.iterrows():
        val0 = str(row[0]).strip() if pd.notna(row[0]) else ""
        val1 = str(row[1]).strip() if pd.notna(row[1]) else ""
        val2 = str(row[2]).strip() if pd.notna(row[2]) else ""
        val3 = str(row[3]).strip() if pd.notna(row[3]) else ""
        val15 = row[15] if pd.notna(row[15]) else 0

        # Store Employee ID & Name header block (matches digits or non-empty ID)
        if val0 and val2 and val0 not in ["Employee Number", "Code"]:
            emp_id, emp_name = val0, val2
            curr_date = None
            continue

        # Forward-fill date across daily entries
        if val1 not in ["", "Date"]:
            curr_date = val1

        # Extract valid line items with non-zero Hours
        if curr_date and val3 not in ["", "Action"] and emp_id:
            try:
                hrs = float(val15)
                if hrs > 0:
                    records.append({
                        "Employee ID": emp_id,
                        "Employee Name": emp_name,
                        "Action": val3,
                        "Date": curr_date,
                        "Hours": hrs,})
            except (ValueError, TypeError):
                continue

    # 3. Handle empty records safely using explicit default columns
    cols = ["Employee ID", "Employee Name", "Action", "Date", "Hours"]
    df = pd.DataFrame(records, columns=cols)

    # Clean & Format Date and Employee ID
    df["Employee ID"] = df["Employee ID"].astype(str)
    df["Date"] = pd.to_datetime(df["Date"]).dt.strftime("%m/%d/%Y")
    df["Hours"] = df["Hours"].round(2)
    df = df[df["Hours"] != 0]

    # Group & Aggregate including Action
    df_aggregated = (
        df.groupby(["Employee ID", "Action", "Date"], as_index=False)
        .agg({"Hours": "sum", "Employee Name": "first"}))

    # Calculate Weekend date (Legacy logic)
    max_date = pd.to_datetime(df_aggregated["Date"], format="%m/%d/%Y").max()
    df_aggregated["Weekend"] = pd.offsets.Week(weekday=5).rollforward(max_date).strftime("%m%d%Y")

    df_aggregated = df_aggregated[["Employee ID", "Employee Name", "Date", "Hours", "Weekend", "Action"]]

    return df_aggregated

#Time and Labor Form

def TimeAndLaborForm(df_customer, query_file):

    # Load data & template
    df_customer = df_customer[df_customer["Action"].astype(str).str.lower() == "lunch"].copy()
    df_query = pd.read_excel(query_file, header=1, dtype=str)
    wb = load_workbook(TEMPLATE_PATH)
    ws = wb["Time & Labor Form"]

    ws["E7"] = datetime.now().strftime("%m/%d/%Y")

    # 2. Clean IDs & Merge (Changed 'inner' to 'left' to keep all df_customer rows)
    query_col = df_query.columns[18]
    df_customer["Employee ID"] = df_customer["Employee ID"].astype(str).str.split(".").str[0].str.strip()
    df_query[query_col] = df_query[query_col].astype(str).str.split(".").str[0].str.strip()

    df_clean = pd.merge(
        df_customer,
        df_query.drop_duplicates(subset=[query_col], keep="first"),
        left_on="Employee ID",
        right_on=query_col,
        how="left")

    # 3. Process & Write Data
    if not df_clean.empty:
        pay_rate = pd.to_numeric(df_clean.get("Current Pay Rate", pd.Series(0, index=df_clean.index)), errors="coerce").fillna(0)
        bill_rate = pd.to_numeric(df_clean.get("Current Bill Rate", pd.Series(0, index=df_clean.index)), errors="coerce").fillna(0)
        order_id = df_clean.get("Order ID", pd.Series("")).fillna("").astype(str).str.split(".").str[0]

        df_tei = pd.DataFrame({
            "A": pd.to_datetime(df_clean["Date"]).dt.strftime("%m/%d/%Y"),
            "B": df_clean["Employee Name_x"],
            "C": df_clean["Employee ID_x"].astype(str).str.split(".").str[0].str.zfill(11),
            "D": order_id.apply(lambda x: x.zfill(11) if x else ""),
            "E": df_clean["Hours"].round(2),
            "G": pay_rate,
            "H": bill_rate,
            "I": (df_clean["Hours"].round(2) * pay_rate).round(2),
            "K": "BILL",
            "L": "MBK"
        })

        for r_idx, row in enumerate(df_tei.itertuples(index=False), start=29):
            for col_name, val in zip(df_tei.columns, row):
                ws[f"{col_name}{r_idx}"] = val

        start_row = 29 + len(df_tei)
    else:
        start_row = 29

    # 4. Clean Extra Rows & Save
    if start_row <= 1000:
        ws.delete_rows(start_row, amount=1000 - start_row + 1)

    output_buffer = BytesIO()
    wb.save(output_buffer)
    output_buffer.seek(0)

    return output_buffer

# CreateUpload

def generate_tei_output(df_customer, query_file):

    df_sick = df_customer[df_customer["Action"].str.lower() == "sick"].reset_index(drop=True)

    df_customer = df_customer[(df_customer["Action"].str.lower() != "sick") & 
                                (df_customer["Action"].str.lower() != "lunch")].reset_index(drop=True)
    df_query = pd.read_excel(query_file, header=1, dtype=str)

    df_customer['Employee ID'] = df_customer['Employee ID'].astype(str).str.split('.').str[0].str.strip()

    query_id_col = df_query.columns[18]
    df_query[query_id_col] = df_query[query_id_col].astype(str).str.split('.').str[0].str.strip()

    df_query_unique = df_query.drop_duplicates(subset=[query_id_col], keep='first')

    df_merged = pd.merge(
        df_customer,
        df_query_unique, 
        left_on='Employee ID', 
        right_on=query_id_col, 
        how='left',
        suffixes=('', '_query'),
        indicator=True)

    error_mask = df_merged['_merge'] == 'left_only'
    df_errors = df_merged.loc[error_mask, ['Employee ID', 'Employee Name', 'Date', 'Hours', 'Weekend']].copy()

    df_clean = df_merged[df_merged['_merge'] == 'both'].copy()

    def process_to_tei(df_input):
        if df_input.empty:
            return pd.DataFrame()
        tei_data = {
            "BRANCH_NUMBER": df_input["Business Unit"].astype(int),
            "PROCESS_DATE": datetime.now().strftime("%m%d%y"),
            "PROCESS_TIME": datetime.now().strftime("%H%M%S"),
            "CUSTOMER_CODE": df_input["Customer ID"].astype(int),
            "TYPE4": "Type4",
            "EMPLOYEE_IDENTIFIER": df_input["Employee ID"].astype(str).str.split('.').str[0].str.zfill(11),
            "ORDER_IDENTIFIER": df_input["Order ID"].astype(str),
            "DAY_DATE": pd.to_datetime(df_input['Date']).dt.strftime('%d%m%Y'),
            "HOURS_DAY": (df_input['Hours'] * 100).round().astype(int),
            "EMPLOYEE_NAME": df_input['Employee Name'],
            "PAY_RATE": 0, "BILL_RATE": 0, "PAY_RATE_OT": 0,
            "BILL_RATE_OT": 0, "PAY_RATE_DT": 0, "BILL_RATE_DT": 0,
            "TIME_RPTG_CD": "", "EARNING_DEDUCTION": "",
            "FLAT_AMOUNT": 0, "MARKUP_PCT": 0, "OTH_EARN_DESCR": "",
            "WEEK_END_DATE": df_input['Weekend']}
        return pd.DataFrame(tei_data)

    df_tei = process_to_tei(df_clean)

    return df_tei, df_errors, df_sick